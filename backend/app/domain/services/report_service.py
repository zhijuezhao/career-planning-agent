from __future__ import annotations

import json
import os
import uuid
from dataclasses import dataclass
from datetime import datetime
from typing import Any

from loguru import logger
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.domain.models.profile_snapshot import ProfileSnapshot
from app.domain.models.report_record import ReportRecord
from app.infrastructure.database import async_session_factory

REPORTS_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "..", "output", "reports")


def _ensure_reports_dir():
    """确保报告输出目录存在。"""
    os.makedirs(REPORTS_DIR, exist_ok=True)


def _safe_json_loads(text: str) -> dict[str, Any]:
    """安全解析 JSON。"""
    text = text.strip()
    if text.startswith("```"):
        lines = text.split("\n")
        text = "\n".join(lines[1:-1] if lines[-1].strip() == "```" else lines[1:])
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        return {}


# ── 新链：records + 惰性 Word ──────────────────────────────────────────────────


async def _latest_snapshot(db: AsyncSession, user_id: int) -> ProfileSnapshot | None:
    """用户最新快照；有 matched_at 的（已匹配）优先。"""
    stmt = (
        select(ProfileSnapshot)
        .where(ProfileSnapshot.user_id == user_id)
        .order_by(
            ProfileSnapshot.matched_at.is_not(None).desc(),
            ProfileSnapshot.created_at.desc(),
        )
        .limit(1)
    )
    return (await db.execute(stmt)).scalar_one_or_none()


async def next_version(user_id: int, db: AsyncSession) -> int:
    """用户下一报告版本号 = 现有记录数 + 1。"""
    cnt = (
        await db.execute(
            select(func.count())
            .select_from(ReportRecord)
            .where(ReportRecord.user_id == user_id)
        )
    ).scalar_one()
    return cnt + 1


async def record_version(record_id: int, user_id: int, db: AsyncSession) -> int:
    """该用户记录按 id 升序的序号，即导出版本号（与创建时 next_version 一致）。"""
    cnt = (
        await db.execute(
            select(func.count())
            .select_from(ReportRecord)
            .where(ReportRecord.user_id == user_id, ReportRecord.id <= record_id)
        )
    ).scalar_one()
    return cnt


async def _build_report_modern(snapshot: ProfileSnapshot, matching_results) -> str:
    """统一 6 模块报告文本：走 resume_agent 的 builder（唯一实现）。"""
    from app.core.resume_agent.tools.report_builder import _build_report

    five = snapshot.five_layers_json or {}
    dims = snapshot.six_dim_scores_json or {}
    basic = (snapshot.form_raw_json or {}).get("basic_info") or {}
    return await _build_report(five, dims, basic, matching_results)


async def create_report_record(
    user_id: int,
    snapshot: ProfileSnapshot,
    matching_results,
    db: AsyncSession,
) -> ReportRecord:
    """落一条 ReportRecord：文本经 _build_report_modern，version = 次版本 + 1。"""
    assert snapshot.user_id == user_id, "快照不属于该用户"
    text = await _build_report_modern(snapshot, matching_results)
    version = await next_version(user_id, db)
    row = ReportRecord(
        user_id=user_id,
        profile_snapshot_id=snapshot.id,
        serial_no=uuid.uuid4(),
        description=f"第{version}版",
        report_text=text,
    )
    db.add(row)
    await db.commit()
    await db.refresh(row)
    logger.info("Report record created | user_id={} | record_id={} | version={}",
                user_id, row.id, version)
    return row


async def ensure_report_word(report: ReportRecord, db: AsyncSession) -> str | None:
    """惰性 Word：已有路径且文件存在 → 直接返回；否则现场生成 + 落库。"""
    if report.word_file_path and os.path.exists(report.word_file_path):
        return report.word_file_path
    _ensure_reports_dir()
    out = os.path.join(REPORTS_DIR, f"{report.serial_no}.docx")
    generate_word_document({"report_text": report.report_text}, out)
    report.word_file_path = out
    await db.commit()
    logger.info("Report word generated | record_id={} | path={}", report.id, out)
    return out


# ── 兼容壳：app/core/agent/tools/report_tool.py 依赖 create_report（保持不编辑）──


@dataclass
class _CompatReportView:
    """create_report 兼容返回视图：暴露工具读取的 5 个属性，不污染 ORM 行。"""
    id: int
    version: int
    target_job: str | None
    word_file_path: str | None
    report_content: dict


async def create_report(
    user_id: int,
    profile_id: int,
    target_job: str | None = None,
    session=None,
) -> _CompatReportView:
    """兼容旧链 create_report：解析用户最新快照 → 落 ReportRecord（空匹配上下文）。

    report_tool.generate_career_report 以 (user_id, profile_id, target_job, session)
    调用并读取 .id/.version/.target_job/.word_file_path/.report_content.get("report_text")。
    """
    own_session = session is None
    if own_session:
        session = async_session_factory()

    try:
        if own_session:
            async with session:
                return await _create_report_compat(user_id, profile_id, target_job, session)
        return await _create_report_compat(user_id, profile_id, target_job, session)
    except Exception as exc:
        logger.warning("Failed to create report | error={}", exc)
        raise


async def _create_report_compat(
    user_id: int,
    profile_id: int,
    target_job: str | None,
    session: AsyncSession,
) -> _CompatReportView:
    snapshot = await _latest_snapshot(session, user_id)
    if snapshot is None:
        raise ValueError("用户能力画像不存在")
    # Tool 路径无实时匹配结果 —— 空列表（报告文本带空匹配上下文生成）
    record = await create_report_record(user_id, snapshot, matching_results=[], db=session)
    return _CompatReportView(
        id=record.id,
        version=await record_version(record.id, user_id, session),
        target_job=target_job,
        word_file_path=record.word_file_path,
        report_content={"report_text": record.report_text},
    )


# ── 兼容壳：旧 5 名字（report_tool 祖父链/旧测试 import 用，永不触已删表）────────


async def get_user_profile_data(
    user_id: int,
    profile_id: int,
    session: AsyncSession,
) -> dict[str, Any] | None:
    """旧链兼容：返回最新快照派生画像；无快照 → None。"""
    snapshot = await _latest_snapshot(session, user_id)
    if snapshot is None:
        return None
    five = snapshot.five_layers_json or {}
    form = snapshot.form_raw_json or {}
    return {
        "direction_tag": (five.get("intention") or {}).get("direction_tag"),
        "intention": five.get("intention") or {},
        "traits": five.get("traits") or {},
        "practice": five.get("practice") or {},
        "soft_skills": five.get("soft_skills") or {},
        "hard_skills": five.get("hard_skills") or {},
        "basic_info": form.get("basic_info") or {},
    }


async def get_dimension_scores_data(
    profile_id: int,
    session: AsyncSession,
) -> dict[str, Any]:
    """旧链兼容：快照六维评分（1:1 约定 profile_id ≈ user）；无 → {}。"""
    snapshot = await _latest_snapshot(session, profile_id)
    if snapshot is None:
        return {}
    return snapshot.six_dim_scores_json or {}


async def get_latest_match_results(
    user_id: int,
    session: AsyncSession,
    top_k: int = 5,
) -> list[dict[str, Any]]:
    """旧链兼容：匹配结果不再落表，新链由 /reports/generate 显式传入 → 空列表。"""
    return []


async def get_latest_career_path(
    user_id: int,
    session: AsyncSession,
) -> dict[str, Any] | None:
    """旧链兼容：从最新快照派生目标岗位；无快照 → None。"""
    snapshot = await _latest_snapshot(session, user_id)
    if snapshot is None:
        return None
    intention = (snapshot.form_raw_json or {}).get("intention") or {}
    targets = intention.get("target_position") or []
    return {
        "target_position": targets[0] if targets else None,
        "path_type": None,
        "milestones": [],
        "learning_resources": {},
        "generated_plan": {},
    }


async def get_latest_growth_plan(
    user_id: int,
    session: AsyncSession,
) -> dict[str, Any] | None:
    """旧链兼容：成长计划不再落表 → 安全空值 None。"""
    return None


# ── Word 导出（原样保留）─────────────────────────────────────────────────────


def generate_word_document(
    report_content: dict[str, Any],
    output_path: str,
) -> str:
    """生成 Word 文档。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    _ensure_reports_dir()

    doc = Document()
    style = doc.styles["Normal"]
    font = style.font
    font.name = "微软雅黑"
    font.size = Pt(11)

    title = doc.add_heading("生涯发展报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    basic_info = report_content.get("basic_info", {})
    if basic_info.get("target_job"):
        doc.add_paragraph(f"目标岗位：{basic_info['target_job']}")
    doc.add_paragraph(f"生成时间：{datetime.utcnow().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph()

    report_text = report_content.get("report_text", "")
    for line in report_text.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith("*") and line.endswith("*"):
            p = doc.add_paragraph()
            run = p.add_run(line.strip("*"))
            run.italic = True
        else:
            doc.add_paragraph(line)

    doc.save(output_path)
    logger.info("Word document generated | path={}", output_path)
    return output_path